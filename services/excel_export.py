from io import BytesIO

import pandas as pd
from openpyxl.styles import Font

from database.repository import Repository


def export_excel(repository: Repository) -> bytes:
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        repository.devices().to_excel(writer, index=False, sheet_name="Dispositivos")
        repository.campaigns().to_excel(writer, index=False, sheet_name="Campanas")
        repository.measurements().to_excel(writer, index=False, sheet_name="Mediciones")
        repository.ztc_results().to_excel(writer, index=False, sheet_name="Resultados ZTC")
        repository.all_points().to_excel(writer, index=False, sheet_name="Puntos V I")
        repository.tracks().to_excel(writer, index=False, sheet_name="Tracks Vt")
        repository.all_track_points().to_excel(writer, index=False, sheet_name="Puntos Track Vt")
        repository.campaign_links().to_excel(writer, index=False, sheet_name="Secuencia Campanas")
        repository.measurement_groups().to_excel(writer, index=False, sheet_name="Grupos")
        group_rows = []
        for group in repository.measurement_groups().itertuples():
            for row in repository.group_measurements(int(group.id)).itertuples():
                group_rows.append({"grupo_id": group.id, "grupo": group.nombre, "tipo": "I-V", "elemento_id": row.id,
                                   "dispositivo": row.dispositivo, "campana": row.campana, "archivo": row.archivo})
            for row in repository.group_tracks(int(group.id)).itertuples():
                group_rows.append({"grupo_id": group.id, "grupo": group.nombre, "tipo": "Track Vt", "elemento_id": row.id,
                                   "dispositivo": row.dispositivo, "campana": row.campana, "archivo": row.archivo})
        pd.DataFrame(group_rows, columns=["grupo_id", "grupo", "tipo", "elemento_id", "dispositivo", "campana", "archivo"]).to_excel(writer, index=False, sheet_name="Elementos grupos")
        repository.unclassified().to_excel(writer, index=False, sheet_name="Sin clasificar")
        workbook = writer.book
        _link_rows(workbook["Mediciones"], "archivo", workbook["Puntos V I"], "medicion_id")
        _link_rows(workbook["Tracks Vt"], "archivo", workbook["Puntos Track Vt"], "track_id")
    return output.getvalue()


def _link_rows(source, source_label: str, target, target_id: str) -> None:
    source_headers = {cell.value: cell.column for cell in source[1]}
    target_headers = {cell.value: cell.column for cell in target[1]}
    targets = {}
    for row in range(2, target.max_row + 1):
        value = target.cell(row, target_headers[target_id]).value
        targets.setdefault(value, row)
    label_column = source_headers.get(source_label)
    id_column = source_headers.get("id")
    if not label_column or not id_column:
        return
    for row in range(2, source.max_row + 1):
        target_row = targets.get(source.cell(row, id_column).value)
        if target_row:
            cell = source.cell(row, label_column)
            cell.hyperlink = f"#'{target.title}'!A{target_row}"
            cell.font = Font(color="0563C1", underline="single")


def export_integrity(repository: Repository) -> dict[str, int]:
    checks = {
        "mediciones_duplicadas": "SELECT COUNT(*) - COUNT(DISTINCT campana_id || ':' || archivo) FROM mediciones",
        "puntos_iv_duplicados": "SELECT COUNT(*) - COUNT(DISTINCT medicion_id || ':' || v || ':' || i) FROM puntos",
        "tracks_duplicados": "SELECT COUNT(*) - COUNT(DISTINCT dispositivo_id || ':' || archivo || ':' || canal) FROM tracks_vt",
        "puntos_track_duplicados": "SELECT COUNT(*) - COUNT(DISTINCT track_id || ':' || t || ':' || vt) FROM puntos_track_vt",
    }
    return {name: max(0, int(repository.connection.execute(sql).fetchone()[0] or 0)) for name, sql in checks.items()}
